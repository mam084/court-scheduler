"""Private lessons: 1-hour sessions a coach books on a court.

Privates are booked on top of the practice schedule (they don't move team
practices). Each one is checked against the schedule: a court that's taken
by a practice, program or another private, or a coach who's busy then, is
flagged so it can be moved.

Privates sheet: coach, date, time (start, e.g. "4:00 PM"), court ("P" =
practice court, the default), type (Team / Personal), team (team privates),
player, note.
"""

from __future__ import annotations

import io
import re
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import DAYS, PRACTICE_COURT, PRIVATE_MINUTES, SLOT_TIMES, court_name, court_short
from .data import blank, parse_courts, to_date

TYPES = ["Team", "Personal"]


def parse_start(value) -> int:
    """'4:00 PM' / '4pm' / '16:30' / '4:30' -> minutes after midnight.
    Without AM/PM, 1-8 o'clock is read as afternoon."""
    t = str(value).strip().lower().replace(".", "").replace(" ", "")
    m = re.match(r"^(\d{1,2})(?::(\d{2}))?(am|pm)?$", t)
    if not m:
        raise ValueError(f"Can't read the time '{value}' (use e.g. 4:00 PM)")
    h, mins, ampm = int(m.group(1)), int(m.group(2) or 0), m.group(3)
    if ampm == "pm" and h < 12:
        h += 12
    elif ampm == "am" and h == 12:
        h = 0
    elif ampm is None and 1 <= h <= 8:
        h += 12
    if not (0 <= h < 24 and 0 <= mins < 60):
        raise ValueError(f"Can't read the time '{value}'")
    return h * 60 + mins


def clock(minutes: int) -> str:
    """Minutes after midnight -> '4:15 PM'."""
    h, m = divmod(int(minutes), 60)
    return f"{h % 12 or 12}:{m:02d} {'AM' if h < 12 else 'PM'}"


def private_type(value) -> str:
    """'Team' or 'Personal' from a cell (anything starting with 't' is a team private)."""
    return "Team" if str(value or "").strip().lower().startswith("t") else "Personal"


def rows(df: pd.DataFrame | None, year: int) -> tuple[list[dict], list[str]]:
    """The Privates sheet as clean dicts (sorted by date and time), plus a
    note for every row that couldn't be read."""
    out, problems = [], []
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return out, problems
    for i, r in df.reset_index(drop=True).iterrows():
        if blank(r.get("coach")) and blank(r.get("date")):
            continue
        where = f"Privates row {i + 2}"
        try:
            if blank(r.get("coach")):
                raise ValueError("needs a coach")
            d = to_date(r["date"], year)
            start = parse_start(r.get("time"))
            courts = parse_courts(r.get("court")) if not blank(r.get("court")) else [PRACTICE_COURT]
            kind = private_type(r.get("type"))
        except (ValueError, TypeError, KeyError) as e:
            problems.append(f"{where}: {e}")
            continue
        team = "" if blank(r.get("team")) else str(r["team"]).strip()
        player = "" if blank(r.get("player")) else str(r["player"]).strip()
        what = (f"Team · {team}" if kind == "Team" else "Personal") + (f" · {player}" if player else "")
        out.append({
            "row": i, "coach": str(r["coach"]).strip(), "date": d, "day": DAYS[d.weekday()],
            "start": start, "end": start + PRIVATE_MINUTES, "time": clock(start),
            "end_label": clock(start + PRIVATE_MINUTES), "court": courts[0], "type": kind, "team": team,
            "player": player, "note": "" if blank(r.get("note")) else str(r["note"]).strip(), "what": what,
        })
    out.sort(key=lambda p: (p["date"], p["start"], p["coach"]))
    return out, problems


def _slot_span(slot: str) -> tuple[int, int]:
    (sh, sm), (eh, em) = SLOT_TIMES[slot]
    return sh * 60 + sm, eh * 60 + em


def conflicts(privs: list[dict], assignments: pd.DataFrame | None,
              closures: dict | None = None) -> dict[int, list[str]]:
    """{private's row: [problems]}: the court is taken (practice, program,
    another private), the coach is busy (a practice or program they coach,
    another private), or the club is closed."""
    out: dict[int, list[str]] = {}
    closures = closures or {}
    a = assignments if assignments is not None and len(assignments) else None
    for i, p in enumerate(privs):
        probs = []
        if p["date"] in closures:
            probs.append(f"the club is closed ({closures[p['date']] or 'closure'})")
        if a is not None:
            day = a[a["date"] == p["date"]]
            for r in day.itertuples():
                if r.slot not in SLOT_TIMES:
                    continue
                s0, s1 = _slot_span(r.slot)
                if not (s0 < p["end"] and p["start"] < s1):
                    continue
                if pd.notna(r.court) and int(r.court) == p["court"]:
                    probs.append(f"{court_name(p['court'])} has {r.name} ({r.slot_label})")
                staff = [c.strip() for c in str(r.coach or "").split(";")]
                if p["coach"] in staff:
                    probs.append(f"{p['coach']} coaches {r.name} then ({r.slot_label})")
        for j, q in enumerate(privs):
            if j == i or q["date"] != p["date"] or not (q["start"] < p["end"] and p["start"] < q["end"]):
                continue
            if q["court"] == p["court"]:
                probs.append(f"another private on {court_name(p['court'])} ({q['coach']}, {q['time']})")
            if q["coach"] == p["coach"]:
                probs.append(f"{p['coach']} has another private then ({q['time']})")
        if probs:
            out[p["row"]] = list(dict.fromkeys(probs))
    return out


def free_courts(when: date, start: int, assignments: pd.DataFrame | None, privs: list[dict], courts) -> list[int]:
    """Courts with nothing booked for the hour starting at `start`."""
    taken = set()
    if assignments is not None and len(assignments):
        for r in assignments[assignments["date"] == when].itertuples():
            if r.slot in SLOT_TIMES and pd.notna(r.court):
                s0, s1 = _slot_span(r.slot)
                if s0 < start + PRIVATE_MINUTES and start < s1:
                    taken.add(int(r.court))
    for q in privs:
        if q["date"] == when and q["start"] < start + PRIVATE_MINUTES and start < q["end"]:
            taken.add(q["court"])
    return [c for c in courts if c not in taken]


_HEAD = PatternFill("solid", fgColor="1F2A44")
_TEAM = PatternFill("solid", fgColor="EEF2FF")


def to_excel(privs: list[dict], coaches: list[str], problems: dict[int, list[str]] | None = None) -> bytes:
    """Every coach (including those with none) and all their privates:
    a summary, one list of everything, and a sheet per coach."""
    problems = problems or {}
    wb = Workbook()
    summary = wb.active
    summary.title = "Coaches"
    every = sorted(set(coaches) | {p["coach"] for p in privs}, key=str.lower)
    summary.append(["Coach", "Team privates", "Personal privates", "Total", "Hours"])
    for c in every:
        mine = [p for p in privs if p["coach"] == c]
        n_team = sum(p["type"] == "Team" for p in mine)
        summary.append([c, n_team, len(mine) - n_team, len(mine), len(mine) * PRIVATE_MINUTES / 60])
    summary.append(["Total", sum(p["type"] == "Team" for p in privs), sum(p["type"] != "Team" for p in privs),
                    len(privs), len(privs) * PRIVATE_MINUTES / 60])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    _style(summary, [24, 14, 17, 9, 9])

    cols = ["Coach", "Date", "Day", "Start", "End", "Court", "Type", "Team", "Player", "Note", "Conflict"]
    allsh = wb.create_sheet("All privates")
    allsh.append(cols)
    for p in sorted(privs, key=lambda p: (p["coach"].lower(), p["date"], p["start"])):
        allsh.append(_line(p, problems))
        if p["type"] == "Team":
            for cell in allsh[allsh.max_row]:
                cell.fill = _TEAM
    _style(allsh, [20, 12, 6, 10, 10, 15, 10, 14, 18, 24, 40])

    used = {"Coaches", "All privates"}
    for c in every:
        mine = [p for p in privs if p["coach"] == c]
        if not mine:
            continue
        name = re.sub(r"[\[\]:*?/\\]", "", c)[:31] or "Coach"
        while name in used:
            name = name[:28] + f" {len(used)}"
        used.add(name)
        sh = wb.create_sheet(name)
        sh.append(cols[1:])
        for p in mine:
            sh.append(_line(p, problems)[1:])
            if p["type"] == "Team":
                for cell in sh[sh.max_row]:
                    cell.fill = _TEAM
        sh.append([])
        sh.append([f"{len(mine)} private{'s' if len(mine) != 1 else ''} · "
                   f"{sum(p['type'] == 'Team' for p in mine)} team, {sum(p['type'] != 'Team' for p in mine)} personal"])
        _style(sh, [12, 6, 10, 10, 15, 10, 14, 18, 24, 40])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _line(p: dict, problems: dict) -> list:
    return [p["coach"], p["date"].strftime("%m/%d/%Y"), p["day"], p["time"], p["end_label"], court_name(p["court"]),
            p["type"], p["team"], p["player"], p["note"], "; ".join(problems.get(p["row"], []))]


def _style(sh, widths):
    for cell in sh[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), _HEAD
        cell.alignment = Alignment(horizontal="center")
    for j, w in enumerate(widths, start=1):
        sh.column_dimensions[get_column_letter(j)].width = w
    sh.freeze_panes = "A2"


def court_choice_label(c) -> str:
    """A court's name for the privates form, noting the practice court is for programs and privates."""
    return court_name(c) + (" (programs & privates)" if c == PRACTICE_COURT else "")


__all__ = ["TYPES", "clock", "conflicts", "court_choice_label", "court_short", "free_courts", "parse_start",
           "private_type", "rows", "to_excel"]
