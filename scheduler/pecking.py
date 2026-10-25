"""The club's pecking order: read it from any spreadsheet and match it to
the teams in a schedule.

A pecking-order sheet is any table with a rank column and a team column,
for example the club's own:

    Rank | Gender | Team Name   | Head Coach | Asst. Coach
    18-1 | Boy    | 18-Black    | Tyler      | Alex
    18-1 | Girl   | 18-Danielle | Danielle   | Vince
    16-1 | Girl   | 16-D&D      | Danielle   | Darci

Column names are matched loosely ("Rank" / "Pecking order" / "Order";
"Team Name" / "Team"; "Head Coach" / "Head"; "Asst. Coach" / "Assistants";
"Gender" / "Side" / "Club"). Teams are matched by name first ("18-Black" =
"18 Black"), then by age + head coach (+ side), so nicknames like "16-D&D"
still find "16-Danielle & Darci".
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd

from .data import blank, parse_names, parse_rank, _team_name_head

_COLS = {
    "rank": ("rank", "pecking", "order", "level"),
    "team": ("team",),
    "head": ("head",),
    "assistants": ("asst", "assistant"),
    "side": ("gender", "side", "club", "boys/girls"),
}
_NO_ONE = {"tba", "tbd", "?", "none", "n/a", "na"}


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _side(value) -> str:
    v = "" if blank(value) else str(value).strip().lower()
    if v.startswith(("boy", "b", "m")) or "seaside" in v:
        return "Boys"
    if v.startswith(("girl", "g", "f", "w")) or "sdvbc" in v:
        return "Girls"
    return ""


def team_side(club, name: str) -> str:
    """Boys/Girls for a team in the app: from its club, else its name
    ('<age> Black/Green' = boys, '<age>-<Coach>' = girls)."""
    s = _side(club)
    if s:
        return s
    low = str(name).lower()
    if re.search(r"\b(black|green)$", low):
        return "Boys"
    if re.match(r"^\s*\d{1,2}\s*[-–]\s*[a-z]", low):
        return "Girls"
    return ""


def _find_header(rows: list[list]) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(rows[:15]):
        labels = [str(v).strip().lower() if v is not None else "" for v in row]
        found = {}
        for key, words in _COLS.items():
            for j, lab in enumerate(labels):
                if lab and any(w in lab for w in words) and j not in found.values():
                    if key == "team" and ("head" in lab or "asst" in lab):
                        continue
                    found[key] = j
                    break
        if "rank" in found and "team" in found:
            return i, found
    return None


def read_pecking(file, filename: str = "") -> pd.DataFrame:
    """Rows of a pecking-order table from an .xlsx (any tab that has Rank
    and Team columns; a tab named like 'pecking' is tried first) or .csv.
    Values are kept raw (Excel dates included) so ranks Excel turned into
    dates can be read back."""
    name = (filename or getattr(file, "name", "") or "").lower()
    data = file.read() if hasattr(file, "read") else open(file, "rb").read()
    tables: list[tuple[str, list[list]]] = []
    if name.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(data), header=None, dtype=str, keep_default_na=False)
        tables.append(("csv", df.values.tolist()))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheets = sorted(wb.worksheets, key=lambda ws: 0 if "peck" in ws.title.lower() or "rank" in ws.title.lower()
                        else 1)
        for ws in sheets:
            rows = [list(r) for r in ws.iter_rows(max_row=400, values_only=True)]
            tables.append((ws.title, rows))
    for title, rows in tables:
        hit = _find_header(rows)
        if not hit:
            continue
        i, cols = hit
        out = []
        for r in rows[i + 1:]:
            get = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
            if blank(get("team")) and blank(get("rank")):
                continue
            out.append({"rank": get("rank"), "team": "" if blank(get("team")) else str(get("team")).strip(),
                        "head": "" if blank(get("head")) else str(get("head")).strip(),
                        "assistants": "" if blank(get("assistants")) else str(get("assistants")).strip(),
                        "side": _side(get("side"))})
        df = pd.DataFrame(out, columns=["rank", "team", "head", "assistants", "side"])
        df.attrs["sheet"] = title
        return df
    raise ValueError("Couldn't find a table with a Rank column and a Team column in that file.")


@dataclass
class Match:
    row: int                      # 0-based row in the pecking table
    sheet_team: str
    rank_text: str                # as it should read in the app, e.g. "16-3"
    rank: int | None
    team: str | None              # matched team in the app
    how: str                      # "name", "age + head coach", "" (not matched)
    note: str = ""
    sheet_assistants: list[str] = field(default_factory=list)


def _sheet_age(rank_value, team_name: str) -> int | None:
    if isinstance(rank_value, (datetime, date, pd.Timestamp)):
        return None  # can't tell month from day without the team; use the name
    m = re.match(r"^\s*(\d{1,2})\s*[-–/ ]\s*\d", str(rank_value or ""))
    if m:
        return int(m.group(1))
    m = re.match(r"^\s*(\d{1,2})", team_name)
    return int(m.group(1)) if m else None


def match_teams(pecking: pd.DataFrame, teams: pd.DataFrame) -> list[Match]:
    """Match every pecking-order row to a team in the app's Teams table."""
    app = []
    for _, t in teams.iterrows():
        name = "" if blank(t.get("team")) else str(t["team"]).strip()
        if not name:
            continue
        try:
            age = int(float(t["age"])) if not blank(t.get("age")) else int(re.match(r"\s*(\d{1,2})", name).group(1))
        except (TypeError, ValueError, AttributeError):
            age = None
        heads = parse_names(t.get("head_coach")) or _team_name_head(name)
        app.append({"name": name, "norm": _norm(name), "age": age, "heads": [h.lower() for h in heads],
                    "side": team_side(t.get("club"), name)})
    used: set[str] = set()
    out = []
    for i, r in pecking.reset_index(drop=True).iterrows():
        sheet_team = r["team"]
        age = _sheet_age(r["rank"], sheet_team)
        if age is None:
            m = re.match(r"^\s*(\d{1,2})", sheet_team)
            age = int(m.group(1)) if m else None
        rank, problem = parse_rank(r["rank"], age)
        rank_text = f"{age}-{rank}" if age and rank else ("" if rank is None else str(rank))
        assts = [a for a in parse_names(r.get("assistants")) if a.lower() not in _NO_ONE]

        team, how = None, ""
        exact = [a for a in app if a["norm"] == _norm(sheet_team) and a["name"] not in used]
        if exact:
            team, how = exact[0]["name"], "name"
        else:
            heads = [h.lower() for h in parse_names(r.get("head"))] or [h.lower() for h in _team_name_head(sheet_team)]
            cands = [a for a in app if a["name"] not in used and a["age"] == age and set(heads) & set(a["heads"])]
            if len(cands) > 1 and r.get("side"):
                same = [a for a in cands if a["side"] == r["side"]]
                cands = same or cands
            if len(cands) == 1:
                team, how = cands[0]["name"], "age + head coach"
            elif len(cands) > 1:
                problem = (problem + "; " if problem else "") + "more than one team fits: " + \
                    ", ".join(a["name"] for a in cands)
        if team:
            used.add(team)
            a_age = next(a["age"] for a in app if a["name"] == team)
            if age and a_age and age != a_age:
                problem = (problem + "; " if problem else "") + f"rank says {age}s but {team} is {a_age}s"
                rank_text = f"{a_age}-{rank}" if rank else rank_text
        elif not problem:
            problem = "no team in the app matches"
        out.append(Match(i, sheet_team, rank_text, rank, team, how, problem or "", assts))
    return out


def unranked(matches: list[Match], teams: pd.DataFrame) -> list[str]:
    got = {m.team for m in matches if m.team}
    return [str(t) for t in teams["team"] if not blank(t) and str(t).strip() not in got]


def apply(teams: pd.DataFrame, matches: list[Match], update_assistants: bool = False) -> pd.DataFrame:
    """Teams table with the matched ranks (and optionally assistants) written in."""
    out = teams.copy()
    if "rank" not in out.columns:
        out["rank"] = ""
    for m in matches:
        if not m.team or m.rank is None:
            continue
        mask = out["team"].astype(str).str.strip() == m.team
        out.loc[mask, "rank"] = m.rank_text
        if update_assistants:
            heads = set()
            for _, t in out[mask].iterrows():
                heads |= set(parse_names(t.get("head_coach")) or _team_name_head(m.team))
            out.loc[mask, "assistants"] = ", ".join(a for a in m.sheet_assistants if a not in heads)
    return out


def assistant_changes(matches: list[Match], teams: pd.DataFrame) -> list[tuple[str, str, str]]:
    """(team, in the app, in the sheet) for each team whose assistants differ."""
    by, heads = {}, {}
    for _, t in teams.iterrows():
        if blank(t.get("team")):
            continue
        name = str(t["team"]).strip()
        by[name] = "" if blank(t.get("assistants")) else str(t["assistants"])
        heads[name] = parse_names(t.get("head_coach")) or _team_name_head(name)
    out = []
    for m in matches:
        if not m.team:
            continue
        cur = [a for a in parse_names(by.get(m.team, "")) if a.lower() not in _NO_ONE]
        sheet = [a for a in m.sheet_assistants if a not in heads.get(m.team, [])]  # a co-head isn't a change
        if sorted(cur) != sorted(sheet):
            out.append((m.team, ", ".join(cur) or "(none)", ", ".join(sheet) or "(none)"))
    return out


def order_table(teams: pd.DataFrame) -> pd.DataFrame:
    """Teams in pecking order (top first), with a 'court priority' number."""
    from .config import pecking_key
    rows = []
    for _, t in teams.iterrows():
        name = "" if blank(t.get("team")) else str(t["team"]).strip()
        if not name:
            continue
        try:
            age = int(float(t["age"])) if not blank(t.get("age")) else int(re.match(r"\s*(\d{1,2})", name).group(1))
        except (TypeError, ValueError, AttributeError):
            continue
        rank, _ = parse_rank(t.get("rank"), age)
        if rank is None:
            low = name.lower()
            rank = 1 if low.endswith("black") else 2 if low.endswith("green") else None
        rows.append({"team": name, "club": "" if blank(t.get("club")) else str(t["club"]), "age": age,
                     "rank": "" if blank(t.get("rank")) else str(t["rank"]),
                     "head_coach": "" if blank(t.get("head_coach")) else str(t["head_coach"]),
                     "assistants": "" if blank(t.get("assistants")) else str(t["assistants"]),
                     "_key": pecking_key(age, rank), "_set": rank is not None})
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    df = df.sort_values(["_key", "club", "team"]).reset_index(drop=True)
    keys = list(dict.fromkeys(df["_key"]))
    df.insert(0, "order", df["_key"].map(lambda k: keys.index(k) + 1))
    return df
