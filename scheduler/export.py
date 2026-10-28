"""Exports: an Excel workbook in the club's grid layout, a calendar in the
club's own week-block layout, and .ics calendar files for a coach, team,
club or everyone (importable into Google Calendar today; a live Google
Calendar sync can replace this later)."""

from __future__ import annotations

import hashlib
import io
from datetime import datetime, timedelta, timezone

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import (
    ALL_COURTS, COURTS, DAYS, PRACTICE_COURT, court_name, court_short, SLOT_KEYS, SLOT_TIMES, SLOTS, WEEK_ORDER,
    WEEKDAYS,
)
from .solver import ScheduleResult, week_start
from .views import _hit

_thin = Side(style="thin", color="C9CED8")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_head = PatternFill("solid", fgColor="1F2A44")
_day = PatternFill("solid", fgColor="EEF1F6")
_closed = PatternFill("solid", fgColor="EDEDED")


def _cell_text(rows: pd.DataFrame) -> str:
    rows = rows.assign(o=rows["slot"].map(SLOT_KEYS.index)).sort_values("o")
    lines = []
    for _, r in rows.iterrows():
        weekday = r.get("day") in WEEKDAYS
        mark = " *" if r["kind"] == "Practice" and weekday and not r.get("on_pattern", True) else ""
        who = f" ({str(r['coach']).replace('; ', ', ')})" if r["coach"] else ""
        tag = " [paired]" if _txt(r.get("paired")) else ""
        tag += " [extra]" if r.get("extra") == True else ""  # noqa: E712
        lines.append(f"{r['slot']} {r['name']}{who}{mark}{tag}")
    return "\n".join(lines)


def _txt(v) -> str:
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) or str(v) in ("nan", "None") else str(v)


def _grid_sheet(wb, title, bookings, day_labels, closed=None, cols=None):
    """Courts as rows, days as columns."""
    ws = wb.create_sheet(title=title[:31])
    closed = closed or {}
    cols = cols or list(WEEKDAYS)
    ws.cell(1, 1, "").fill = _head
    for j, wd in enumerate(cols, start=2):
        cell = ws.cell(1, j, day_labels.get(wd, wd))
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), _head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32 if day_labels else 18
    for i, c in enumerate(ALL_COURTS, start=2):
        rc = ws.cell(i, 1, court_name(c))
        rc.font, rc.fill, rc.border = Font(bold=True), _day, _border
        rc.alignment = Alignment(horizontal="center", vertical="center")
        court_rows = bookings[bookings["court"] == c] if len(bookings) else bookings
        n_lines = 1
        for j, wd in enumerate(cols, start=2):
            if wd in closed:
                cell = ws.cell(i, j, closed[wd] if i == 2 else "")
                cell.fill = _closed
            else:
                rows = court_rows[court_rows["day"] == wd] if len(court_rows) else court_rows
                txt = _cell_text(rows) if len(rows) else ""
                cell = ws.cell(i, j, txt)
                n_lines = max(n_lines, txt.count("\n") + 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _border
        ws.row_dimensions[i].height = max(30, 15 * n_lines + 4)
    ws.column_dimensions["A"].width = 14
    for j in range(2, len(cols) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 30
    ws.freeze_panes = "B2"
    return ws


def to_excel(result: ScheduleResult, closures: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    a = result.assignments

    # Recurring template = each team's pattern + the weekly locals.
    tmpl = result.pattern.rename(columns={"team": "name"}).assign(kind="Practice", on_pattern=True)
    locals_ = a[a["kind"] == "Local"].drop_duplicates(subset=["day", "slot", "court", "name"])
    tmpl = pd.concat([tmpl, locals_[tmpl.columns.intersection(locals_.columns)]], ignore_index=True)
    _grid_sheet(wb, "Weekly Template", tmpl, {})
    _calendar_sheet(wb, result, closures)

    weekend_open = set(result.weekend_dates)
    for ws_date, g in a.groupby(a["date"].map(week_start)):
        labels, closed, cols = {}, {}, []
        for i, wd in enumerate(WEEK_ORDER):
            d = ws_date + timedelta(days=i)
            if wd in ("Sat", "Sun") and d not in weekend_open and not any(
                    tn.kind == "local" and tn.covers(d) for tn in result.tournaments):
                continue
            cols.append(wd)
            labels[wd] = f"{wd}\n{d:%b %-d}"
            if d in closures:
                closed[wd] = f"Closed: {closures[d]}"
            elif not (result.start <= d <= result.end):
                closed[wd] = "Outside this schedule period"
            else:
                here = [tn.name + (" (local, until 7 PM)" if tn.kind == "local" else "")
                        for tn in result.tournaments if tn.covers(d)]
                if here:
                    labels[wd] += "\n" + "; ".join(here)
        _grid_sheet(wb, f"Week of {ws_date:%b %-d}", g, labels, closed, cols)

    flat = a.copy()
    flat["date"] = pd.to_datetime(flat["date"]).dt.strftime("%Y-%m-%d")
    flat["court"] = flat["court"].map(court_name)
    flat["note"] = [("Extra practice. " if e == True else "") + _txt(p)  # noqa: E712
                    for e, p in zip(flat.get("extra", [None] * len(flat)), flat.get("paired", [None] * len(flat)))]
    flat = flat[["date", "day", "slot_label", "court", "kind", "name", "club", "coach", "on_pattern", "note"]]
    flat.columns = ["Date", "Day", "Time", "Court", "Type", "Team / Program", "Club", "Coach(es)", "On usual schedule",
                    "Note"]
    sh = wb.create_sheet("All Bookings")
    sh.append(list(flat.columns))
    for row in flat.itertuples(index=False):
        sh.append(list(row))
    for j, w in enumerate([12, 6, 14, 7, 9, 22, 10, 24, 18, 50], start=1):
        sh.column_dimensions[get_column_letter(j)].width = w
    for c in sh[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _head

    if result.tournaments:
        sh = wb.create_sheet("Tournaments")
        sh.append(["Tournament", "Type", "Start", "End", "Teams / courts"])
        for tn in result.tournaments:
            sh.append([tn.name, tn.kind.title(), f"{tn.start:%Y-%m-%d}", f"{tn.end:%Y-%m-%d}",
                       ", ".join(tn.teams) if tn.kind == "travel" else ", ".join(court_name(c) for c in tn.courts)])
        for j, w in enumerate([28, 9, 12, 12, 50], start=1):
            sh.column_dimensions[get_column_letter(j)].width = w
        for c in sh[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), _head

    pc = result.practice_counts
    if pc is not None and len(pc):
        sh = wb.create_sheet("Practice Counts", 1)  # right after the template
        active = pc.groupby("week_of")[["target", "scheduled"]].sum()
        weeks = sorted(active[(active["target"] > 0) | (active["scheduled"] > 0)].index)
        sh.append(["Team"] + [f"Week of {w:%b %-d}" for w in weeks] + ["Total short"])
        red, pink, green, grey = (PatternFill("solid", fgColor=c) for c in ("F8B4B4", "FDE2E1", "EAF6EC", "F4F5F7"))
        lookup = {(r.team, r.week_of): r for r in pc.itertuples()}
        totals = pc.groupby("team")["short"].sum()
        order = list(totals[totals > 0].sort_values(ascending=False).index) + sorted(totals[totals == 0].index)
        for i, tm in enumerate(order, start=2):
            sh.cell(i, 1, tm).font = Font(bold=True)
            for j, w in enumerate(weeks, start=2):
                r = lookup.get((tm, w))
                if r is None or r.target == 0:
                    cell = sh.cell(i, j, r.scheduled if r is not None and r.scheduled else "–")
                    cell.fill = grey
                else:
                    cell = sh.cell(i, j, f"{r.scheduled} of {r.target}")
                    cell.fill = (red if r.short >= 2 else pink) if r.short else green
                    if r.short:
                        cell.font = Font(bold=True, color="9B1C1C")
                cell.alignment = Alignment(horizontal="center")
                cell.border = _border
            tot = sh.cell(i, len(weeks) + 2, int(totals[tm]))
            if totals[tm]:
                tot.font = Font(bold=True, color="9B1C1C")
        sh.column_dimensions["A"].width = 22
        for j in range(2, len(weeks) + 3):
            sh.column_dimensions[get_column_letter(j)].width = 14
        for c in sh[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), _head
        sh.freeze_panes = "B2"

    if len(result.issues):
        sh = wb.create_sheet("Review")
        cols = ["severity", "team", "week_of", "issue", "detail"]
        sh.append([c.replace("_", " ").title() for c in cols])
        for row in result.issues[cols].itertuples(index=False):
            sh.append([str(v) for v in row])
        for j, w in enumerate([10, 14, 12, 34, 50], start=1):
            sh.column_dimensions[get_column_letter(j)].width = w
        for c in sh[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), _head

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ics_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def to_ics(result: ScheduleResult, coach: str, club_name: str = "Club Practice",
           tz: str = "America/Los_Angeles") -> bytes:
    """Calendar file of one coach's practices and locals."""
    a = result.assignments
    mine = a[a["coach"].fillna("").map(lambda s: coach in [x.strip() for x in s.split(";")])]
    return _ics(a, mine, f"{club_name} – {coach}", tz, key_extra=coach)


def calendar_ics(result: ScheduleResult, highlight: set[str] | None = None, include_locals: bool = True,
                 title: str = "Club Practice", tz: str = "America/Los_Angeles") -> bytes:
    """Calendar file of what the Calendar shows: everything, or only the
    picked clubs / teams / coaches. One event per practice or program."""
    a = result.assignments
    sel = a if include_locals else a[a["kind"] == "Practice"]
    if highlight:
        sel = sel[sel.apply(lambda r: _hit(r, highlight), axis=1)]
    name = title + (" – " + ", ".join(sorted(highlight)) if highlight else "")
    return _ics(a, sel, name, tz, key_extra="|".join(sorted(highlight or [])))


def _ics(a: pd.DataFrame, mine: pd.DataFrame, cal_name: str, tz: str, key_extra: str = "") -> bytes:
    mine = mine.drop_duplicates(subset=["date", "slot", "name"])  # locals span several courts
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//court-scheduler//EN",
             f"X-WR-CALNAME:{_ics_escape(cal_name)}", f"X-WR-TIMEZONE:{tz}"]
    coach = key_extra
    for _, r in mine.iterrows():
        (sh, sm), (eh, em) = SLOT_TIMES[r["slot"]]
        d = r["date"]
        dstart = datetime(d.year, d.month, d.day, sh, sm).strftime("%Y%m%dT%H%M%S")
        dend = datetime(d.year, d.month, d.day, eh, em).strftime("%Y%m%dT%H%M%S")
        courts = a[(a["date"] == d) & (a["slot"] == r["slot"]) & (a["name"] == r["name"])]["court"].tolist()
        where = _courts_text(courts, long=True)
        key = f"{d:%Y%m%d}|{r['slot']}|{r['name']}|{coach}"
        uid = hashlib.md5(key.encode()).hexdigest() + "@court-scheduler"
        lines += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{stamp}",
                  f"DTSTART;TZID={tz}:{dstart}", f"DTEND;TZID={tz}:{dend}",
                  f"SUMMARY:{_ics_escape(r['name'])} ({r['kind'].lower()})",
                  f"LOCATION:{_ics_escape(where)}",
                  f"DESCRIPTION:{_ics_escape(SLOTS[r['slot']] + ' · ' + where + (' · ' + r['coach'] if isinstance(r.get('coach'), str) and r['coach'] else ''))}",
                  "END:VEVENT"]
    lines.append("END:VCALENDAR")
    return ("\r\n".join(lines) + "\r\n").encode()


def _courts_text(courts, long: bool = False) -> str:
    """[1, 2, 3, 7] -> '1-3, 7' (as text, so Excel never turns '1-3' into a date)."""
    nums = sorted({int(c) for c in courts if pd.notna(c)})
    parts, run = [], []
    for c in nums + [None]:
        if c is not None and c != PRACTICE_COURT and run and c == run[-1] + 1:
            run.append(c)
            continue
        if run:
            parts.append(court_short(run[0]) if len(run) == 1 else f"{run[0]}-{run[-1]}")
        run = [c] if c is not None else []
    txt = ", ".join(parts)
    if not long:
        return txt
    if nums == [PRACTICE_COURT]:
        return "Practice court"
    return ("Courts " if len(nums) > 1 else "Court ") + txt


def _calendar_sheet(wb, result: ScheduleResult, closures: dict, highlight=None, include_locals: bool = True,
                    title: str = "Calendar"):
    """The club's calendar layout: one block per week (Sunday-Saturday), each
    day three columns: time · what · courts."""
    a = result.assignments
    sel = a if include_locals else a[a["kind"] == "Practice"]
    if highlight:
        sel = sel[sel.apply(lambda r: _hit(r, highlight), axis=1)]
    if len(sel):
        sel = (sel.groupby(["date", "slot", "slot_label", "kind", "name"], as_index=False)
               .agg(court=("court", list), club=("club", "first")))
    by_date = {d: g for d, g in sel.groupby("date")} if len(sel) else {}
    weekend_open = set(result.weekend_dates)
    sh = wb.create_sheet(title)
    bold_w = Font(bold=True, color="FFFFFF")
    local_fill = PatternFill("solid", fgColor="F3F4F6")
    tourn_fill = PatternFill("solid", fgColor="FFF4D6")
    row = 1
    ws = week_start(result.start)
    while ws <= result.end:
        days = [ws + timedelta(days=i) for i in range(7)]
        for j, d in enumerate(days):
            col = 1 + 3 * j
            for k, v in enumerate((d, f"{d:%A}", "Cts")):
                c = sh.cell(row, col + k, v)
                c.font, c.fill, c.border = bold_w, _head, _border
                c.alignment = Alignment(horizontal="center")
            sh.cell(row, col).number_format = "mmm d"
        entries = {}
        for j, d in enumerate(days):
            items = []
            if not (result.start <= d <= result.end):
                items.append(("", "", "", None))
            elif d in closures:
                items.append(("", f"Closed: {closures[d]}", "", _closed))
            else:
                for tn in result.tournaments:
                    if tn.covers(d) and (not highlight or tn.kind == "local" or set(tn.teams) & set(highlight)):
                        where = _courts_text(tn.courts) if tn.kind == "local" else "away"
                        items.append(("until 7 PM" if tn.kind == "local" else "", tn.name, where, tourn_fill))
                if d.weekday() >= 5 and d not in weekend_open and d not in by_date and not items:
                    items.append(("", "Closed", "", _closed))
                g = by_date.get(d)
                if g is not None:
                    g = g.assign(o=g["slot"].map(SLOT_KEYS.index)).sort_values(["o", "name"])
                    for r in g.itertuples():
                        items.append((r.slot_label, r.name, _courts_text(r.court),
                                      local_fill if r.kind == "Local" else None))
            entries[j] = items
        n = max([len(v) for v in entries.values()] + [1])
        for j, items in entries.items():
            col = 1 + 3 * j
            for i in range(n):
                vals = items[i] if i < len(items) else ("", "", "", None)
                for k in range(3):
                    c = sh.cell(row + 1 + i, col + k, vals[k] if vals[k] != "" else None)
                    c.number_format = "@"   # text: '5-7' stays '5-7'
                    c.border = _border
                    c.alignment = Alignment(vertical="top", wrap_text=k == 1)
                    if vals[3] is not None:
                        c.fill = vals[3]
        row += n + 2
        ws += timedelta(days=7)
    for j in range(7):
        sh.column_dimensions[get_column_letter(1 + 3 * j)].width = 11
        sh.column_dimensions[get_column_letter(2 + 3 * j)].width = 18
        sh.column_dimensions[get_column_letter(3 + 3 * j)].width = 6
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToHeight = 0
    return sh


def calendar_excel(result: ScheduleResult, closures: dict, highlight: set[str] | None = None,
                   include_locals: bool = True) -> bytes:
    """Just the calendar, in the club's week-block layout."""
    wb = Workbook()
    wb.remove(wb.active)
    _calendar_sheet(wb, result, closures, highlight, include_locals)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------- team dates (per team) --
# One row per team, one column per date: each practice (time and court) and
# each travel tournament the team goes to, with the tournament days circled.
# The circles are drawn shapes (Excel, Numbers, LibreOffice); the cells are
# also red-outlined, which is what shows in Google Sheets (it drops shapes).

_TOURN_FILL = PatternFill("solid", fgColor="FDECEC")
_TOURN_SIDE = Side(style="medium", color="D62828")
_WEEKEND_FILL = PatternFill("solid", fgColor="F4F6FA")
_DATE_COL_W = 11      # characters
_TEAM_ROW_H = 38      # points


def _slot_text(slot: str) -> str:
    return slot.replace("AM", " AM").replace("PM", " PM")


def team_dates_excel(result: ScheduleResult, closures: dict, order: list[str] | None = None,
                     title: str = "", sections: dict[str, str] | None = None) -> bytes:
    """Every team's practice and tournament dates, one row per team (in
    `order`, e.g. girls then boys, each in pecking order), with tournament
    days circled. `sections` (team -> label, e.g. "Girls") puts a heading
    row wherever the label changes."""
    start, end = result.start, result.end
    a = result.assignments
    pr = a[a["kind"] == "Practice"] if len(a) else a
    teams = list(order or [])
    for t in list(result.pattern["team"]) if len(result.pattern) else []:
        if t not in teams:
            teams.append(t)
    for t in sorted(set(pr["name"])) if len(pr) else []:
        if t not in teams:
            teams.append(t)
    heads = {}
    if len(result.pattern):
        heads = dict(zip(result.pattern["team"], result.pattern["head"]))
    if len(pr):
        for t, h in zip(pr["name"], pr["head"]):
            heads.setdefault(t, h)
    dates = list(pd.date_range(start, end).date)
    travel = [tn for tn in result.tournaments if tn.kind == "travel"]
    locals_t = [tn for tn in result.tournaments if tn.kind == "local"]

    wb = Workbook()
    sh = wb.active
    sh.title = "Team Dates"
    first_col = 3                 # A team, B head coach, then the dates
    head_row, local_row = 3, 4
    top = 5                       # first team row
    sh.cell(1, 1, title or "Team practices & tournaments").font = Font(bold=True, size=14)
    sh.cell(2, 1, f"{start:%b %-d} – {end:%b %-d, %Y}  ·  practice start time and court  ·  "
                  "circled = tournament").font = Font(color="5B6475", italic=True)

    for j, label in ((1, "Team"), (2, "Head coach")):
        c = sh.cell(head_row, j, label)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), _head, Alignment(vertical="center")
    c = sh.cell(local_row, 1, "Club events")
    c.font, c.alignment = Font(italic=True, color="5B6475"), Alignment(vertical="center")
    for k, d in enumerate(dates):
        col = first_col + k
        c = sh.cell(head_row, col, f"{d:%a}\n{d:%-m/%-d}")
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ev = [f"Closed: {closures[d]}"] if d in closures else []
        ev += [f"{tn.name} (local)" for tn in locals_t if tn.covers(d)]
        if ev:
            c = sh.cell(local_row, col, "\n".join(ev))
            c.font = Font(size=8, italic=True, color="5B6475")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sh.column_dimensions[get_column_letter(col)].width = _DATE_COL_W
    total_col = first_col + len(dates)
    for j, label in ((total_col, "Practices"), (total_col + 1, "Tournaments")):
        c = sh.cell(head_row, j, label)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sh.row_dimensions[head_row].height = 30
    sh.row_dimensions[local_row].height = 30 if (closures or locals_t) else 16

    by_team_day = {}
    for r in pr.itertuples() if len(pr) else []:
        by_team_day.setdefault((r.name, r.date), []).append(r)
    circles = []                  # (row, col), 1-based
    sections = sections or {}
    row, last_sec = top - 1, None
    for t in teams:
        sec = sections.get(t)
        if sec and sec != last_sec:
            row += 1
            c = sh.cell(row, 1, sec)
            c.font = Font(bold=True, size=12, color="1F2A44")
            sh.row_dimensions[row].height = 20
            for col in range(1, total_col + 2):
                sh.cell(row, col).fill = PatternFill("solid", fgColor="DCE3EE")
            last_sec = sec
        row += 1
        sh.row_dimensions[row].height = _TEAM_ROW_H
        c = sh.cell(row, 1, t)
        c.font, c.fill, c.border = Font(bold=True), _day, _border
        c.alignment = Alignment(vertical="center")
        c = sh.cell(row, 2, str(heads.get(t, "") or "").replace("; ", ", "))
        c.border, c.alignment = _border, Alignment(vertical="center", wrap_text=True)
        n_prac, my_t = 0, []
        for k, d in enumerate(dates):
            col = first_col + k
            lines = []
            for p in sorted(by_team_day.get((t, d), []), key=lambda p: SLOT_KEYS.index(p.slot)):
                n_prac += 1
                tag = " +" if p.extra == True else ""  # noqa: E712
                lines.append(f"{_slot_text(p.slot)}{tag}\nCt {court_short(p.court)}")
            tn_here = [tn for tn in travel if t in tn.teams and tn.covers(d)]
            for tn in tn_here:
                lines.append(tn.name)
                if tn.name not in my_t:
                    my_t.append(tn.name)
            c = sh.cell(row, col, "\n".join(lines))
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = _border
            if tn_here:
                c.font, c.fill = Font(bold=True, color="B21E1E", size=8), _TOURN_FILL
                c.border = Border(left=_TOURN_SIDE, right=_TOURN_SIDE, top=_TOURN_SIDE, bottom=_TOURN_SIDE)
                circles.append((row, col))
            elif d in closures:
                c.fill = _closed
            elif d.weekday() >= 5:
                c.fill = _WEEKEND_FILL
            if lines and not tn_here:
                c.font = Font(size=10)
        c = sh.cell(row, total_col, n_prac)
        c.alignment, c.border, c.font = Alignment(horizontal="center", vertical="center"), _border, Font(bold=True)
        c = sh.cell(row, total_col + 1, ", ".join(my_t))
        c.alignment, c.border = Alignment(vertical="center", wrap_text=True), _border
        if my_t:
            c.font = Font(bold=True, color="B21E1E")

    sh.column_dimensions["A"].width = 14
    sh.column_dimensions["B"].width = 14
    sh.column_dimensions[get_column_letter(total_col)].width = 10
    sh.column_dimensions[get_column_letter(total_col + 1)].width = 26
    sh.freeze_panes = sh.cell(top, first_col)
    sh.print_title_rows = f"{head_row}:{local_row}"
    sh.print_title_cols = "A:B"
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_margins.left = sh.page_margins.right = 0.3
    sh.page_margins.top = sh.page_margins.bottom = 0.4

    # A plain list too, one row per practice or tournament day (sortable).
    ls = wb.create_sheet("List")
    ls.append(["Team", "Date", "Day", "What", "Time", "Court"])
    for c in ls[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _head
    red = Font(bold=True, color="B21E1E")
    for t in teams:
        for d in dates:
            for p in sorted(by_team_day.get((t, d), []), key=lambda p: SLOT_KEYS.index(p.slot)):
                ls.append([t, d, f"{d:%a}", "Extra practice" if p.extra == True else "Practice",  # noqa: E712
                           p.slot_label, court_name(p.court)])
                ls.cell(ls.max_row, 2).number_format = "m/d/yyyy"
            for tn in travel:
                if t in tn.teams and tn.covers(d):
                    ls.append([t, d, f"{d:%a}", f"Tournament: {tn.name}", "", ""])
                    ls.cell(ls.max_row, 2).number_format = "m/d/yyyy"
                    for c in ls[ls.max_row]:
                        c.font = red
    for j, w in enumerate([14, 11, 6, 34, 12, 14], start=1):
        ls.column_dimensions[get_column_letter(j)].width = w
    ls.freeze_panes = "A2"
    ls.auto_filter.ref = f"A1:F{max(ls.max_row, 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    return _add_circles(buf.getvalue(), "xl/worksheets/sheet1.xml", circles,
                        col_w={first_col + k: _DATE_COL_W for k in range(len(dates))}, row_h=_TEAM_ROW_H)


def _add_circles(xlsx: bytes, sheet_path: str, cells: list[tuple[int, int]], col_w: dict[int, float],
                 row_h: float) -> bytes:
    """Draws a red ellipse inside each (row, col) cell (1-based) of a sheet,
    by adding a drawing part to the saved workbook (openpyxl can't make
    shapes). The ellipses move and resize with their cells."""
    import re
    import zipfile
    if not cells:
        return xlsx
    emu_px, emu_pt = 9525, 12700
    parts = []
    for n, (r, c) in enumerate(cells, start=1):
        w = int((col_w.get(c, 9) * 7 + 5) * emu_px)
        h = int(row_h * emu_pt)
        ix, iy = 3 * emu_px, 2 * emu_px
        parts.append(
            f'<xdr:twoCellAnchor editAs="twoCell">'
            f'<xdr:from><xdr:col>{c - 1}</xdr:col><xdr:colOff>{ix}</xdr:colOff>'
            f'<xdr:row>{r - 1}</xdr:row><xdr:rowOff>{iy}</xdr:rowOff></xdr:from>'
            f'<xdr:to><xdr:col>{c - 1}</xdr:col><xdr:colOff>{w - ix}</xdr:colOff>'
            f'<xdr:row>{r - 1}</xdr:row><xdr:rowOff>{h - iy}</xdr:rowOff></xdr:to>'
            f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{n + 1}" name="Tournament circle {n}"/>'
            f'<xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w - 2 * ix}" '
            f'cy="{h - 2 * iy}"/></a:xfrm><a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom><a:noFill/>'
            f'<a:ln w="25400"><a:solidFill><a:srgbClr val="D62828"/></a:solidFill></a:ln></xdr:spPr>'
            f'</xdr:sp><xdr:clientData fPrintsWithSheet="1"/></xdr:twoCellAnchor>')
    drawing = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">' + "".join(parts) + "</xdr:wsDr>")
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    rels_path = sheet_path.replace("worksheets/", "worksheets/_rels/") + ".rels"
    rid = "rIdCircles1"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        names = src.namelist()
        for name in names:
            data = src.read(name)
            if name == sheet_path:
                xml = data.decode("utf-8")
                tag = (f'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                       f'r:id="{rid}"/>')
                # <drawing> goes after the page setup elements, before any table parts.
                m = re.search(r"<(legacyDrawing|tableParts|extLst)\b", xml)
                xml = xml[:m.start()] + tag + xml[m.start():] if m else xml.replace("</worksheet>", tag + "</worksheet>")
                data = xml.encode("utf-8")
            elif name == "[Content_Types].xml":
                xml = data.decode("utf-8").replace(
                    "</Types>", '<Override PartName="/xl/drawings/drawingCircles1.xml" '
                                'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
                data = xml.encode("utf-8")
            elif name == rels_path:
                xml = data.decode("utf-8").replace(
                    "</Relationships>", f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/'
                                        'officeDocument/2006/relationships/drawing" '
                                        'Target="../drawings/drawingCircles1.xml"/></Relationships>')
                data = xml.encode("utf-8")
            z.writestr(name, data)
        if rels_path not in names:
            z.writestr(rels_path, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                                  '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                  f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/'
                                  '2006/relationships/drawing" Target="../drawings/drawingCircles1.xml"/>'
                                  '</Relationships>')
        z.writestr("xl/drawings/drawingCircles1.xml", drawing)
    return out.getvalue()
